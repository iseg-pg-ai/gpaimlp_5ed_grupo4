"""Import the complementary workbook without promoting historical proposals to inventory."""
from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any
from openpyxl import load_workbook

BASES = {"Atracoes_base": ("atracoes", "id"), "Restaurantes_base": ("restaurantes", "id_blu"), "Experiencias_base": ("experiencias", "nome_da_experiencia"), "Precos_base": ("precos", "id")}
KEYS = {"Roteiros": "id_roteiro", "Viajantes": "id_grupo", "Atividades": "id_atividade", "Experiencias": "id_experiencia", "Precos_propostas": "id_preco", "Fontes_PDF": "codigo_pdf", "Casos_interacoes": "id_caso", "Regras_BLU": "id"}


def integrate_structured(path: Path, tables: dict[str, list[dict[str, Any]]]) -> dict[str, Any]:
    from etl.pipeline import json_value, records_from_section, slug
    workbook = load_workbook(path, read_only=True, data_only=True)
    imported = {}
    try:
        for sheet in workbook:
            name = "structured_" + slug(sheet.title)
            if sheet.title in ("Guia_dataset", "Auditoria_vazios"):
                records = [{"cells": [json_value(v) for v in row], "_source_sheet": sheet.title, "_source_row": n}
                           for n, row in enumerate(sheet.iter_rows(values_only=True), 1) if any(v is not None for v in row)]
            else:
                records = records_from_section(sheet, 5 if sheet.title == "Fontes_enriquecimento" else 1)
            for record in records:
                record["_source_file"] = path.as_posix()
            imported[sheet.title] = records
            tables[name] = records
    finally:
        workbook.close()
    required = set(BASES) | set(KEYS)
    if required - imported.keys():
        raise ValueError(f"Missing structured sheets: {sorted(required - imported.keys())}")
    for sheet, key in {**KEYS, **{s: spec[1] for s, spec in BASES.items()}}.items():
        values = [r.get(key) for r in imported[sheet]]
        if any(v is None for v in values) or len(values) != len(set(values)):
            raise ValueError(f"Missing or duplicate {key} in {sheet}")
    route_ids = {r["id_roteiro"] for r in imported["Roteiros"]}
    pdf_ids = {r["codigo_pdf"] for r in imported["Fontes_PDF"]}
    attraction_ids = {r["id"] for r in imported["Atracoes_base"]}
    relations = []
    for sheet, rows in imported.items():
        for row in rows:
            for field, keys in (("id_roteiro", route_ids), ("codigo_fonte", pdf_ids), ("id_atracao_base", attraction_ids)):
                if field == "codigo_fonte" and not str(row.get(field, "")).startswith("PDF-"):
                    continue  # Explicit non-PDF provenance (e.g. interaction workbook).
                if row.get(field) is not None and row[field] not in keys:
                    relations.append({"sheet": sheet, "row": row["_source_row"], "field": field, "value": row[field]})
    if relations:
        raise ValueError(f"Broken structured references: {relations[:5]}")
    changes, conflicts = [], []
    for sheet, (table, key) in BASES.items():
        index = {r[key]: r for r in tables[table]}
        for record in imported[sheet]:
            identifier = record[key]
            if identifier not in index:
                target = dict(record)
                tables[table].append(target)
                index[identifier] = target
                changes.append({"table": table, "record_id": identifier, "field": "*", "action": "added"})
            else:
                target = index[identifier]
                for field, value in record.items():
                    if field.startswith("_") or field == key or value is None:
                        continue
                    if target.get(field) is None:
                        target[field] = value
                        changes.append({"table": table, "record_id": identifier, "field": field, "action": "filled"})
                    elif target[field] != value:
                        conflicts.append({"table": table, "record_id": identifier, "field": field, "retained_value": target[field], "supplement_value": value, "_source_sheet": sheet, "_source_row": record["_source_row"], "_source_file": path.as_posix()})
            target["_supplement_source"] = {"file": path.as_posix(), "sheet": sheet, "row": record["_source_row"], "status": record.get("estado_de_utilizacao")}
    tables["structured_merge_changes"] = changes
    tables["structured_merge_conflicts"] = conflicts
    return {"file": path.as_posix(), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "sheets": {k: len(v) for k, v in imported.items()}, "changes": len(changes), "conflicts": len(conflicts), "policy": "Fill missing fields by exact key; retain existing values and audit conflicts. Historical proposals remain separate."}
