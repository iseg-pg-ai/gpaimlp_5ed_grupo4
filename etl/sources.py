"""Source discovery boundary. Local folders supply the same files a future S3 adapter can stage."""
from dataclasses import dataclass
import hashlib
from pathlib import Path
from openpyxl import load_workbook

@dataclass(frozen=True)
class LocalSource:
    root: Path

    def discover(self):
        if not self.root.is_dir():
            raise FileNotFoundError(f"Source directory not found: {self.root}")
        main, supplements, pdfs, inventory = [], [], [], []
        for path in sorted(self.root.rglob("*")):
            if not path.is_file() or path.name.startswith("~$") or path.suffix.lower() not in (".xlsx", ".pdf"):
                continue
            if not path.resolve().is_relative_to(self.root.resolve()):
                raise ValueError(f"Source escapes local directory: {path}")
            role = "proposal_pdf"
            if path.suffix.lower() == ".pdf":
                pdfs.append(path)
            else:
                workbook = load_workbook(path, read_only=True, data_only=True)
                try:
                    sheets = set(workbook.sheetnames)
                finally:
                    workbook.close()
                if "REGRAS_CURADORIA" in sheets and "ATRACOES" in sheets:
                    main.append(path)
                    role = "catalog_workbook"
                elif {"Roteiros", "Atracoes_base", "Precos_propostas"} <= sheets:
                    supplements.append(path)
                    role = "structured_supplement"
                else:
                    role = "reference_workbook"
            inventory.append({"key": path.relative_to(self.root).as_posix(), "role": role, "size_bytes": path.stat().st_size, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
        if len(main) != 1 or len(supplements) > 1:
            raise ValueError(f"Expected one main workbook and at most one supplement; found {len(main)} and {len(supplements)}")
        return main[0], supplements[0] if supplements else None, pdfs, {"type": "local", "root": self.root.as_posix(), "objects": inventory}
