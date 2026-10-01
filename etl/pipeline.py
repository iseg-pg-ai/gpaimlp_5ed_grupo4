"""Extract, transform and load all files placed in the data directory.

The pipeline deliberately keeps the source payload (including PDF page text) while
also creating clean relational records.  That makes each produced field auditable
back to a filename, sheet and row/page number.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sqlite3
import unicodedata
import uuid
from collections import Counter
from contextlib import closing
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from pypdf import PdfReader

from etl.analytics import build_analytics
from etl.validation import validate
from etl.sources import LocalSource


# Map each regular worksheet to its header row; the rows above are document titles.
HEADER_ROWS = {
    "DATA_REQUEST": 6,
    "ATRACOES": 6,
    "RESTAURANTES": 1,
    "EXPERIENCIAS": 1,
    "PREÇOS": 1,
}
# Patterns used to make useful, searchable document attributes from PDF text.
EMAIL = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)
PHONE = re.compile(r"(?<!\d)(?:\+351\s*)?(?:\d[\s.-]?){9}(?!\d)")
URL = re.compile(r"https?://[^\s)>]+|www\.[^\s)>]+", re.I)
MONEY = re.compile(r"(?:€\s?\d[\d.,]*|\d[\d.,]*\s?(?:EUR|euros?))", re.I)


def slug(value: str) -> str:
    """Convert a human label into a stable ASCII table/column identifier."""
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", normalized.lower())).strip("_")


def json_value(value: Any) -> Any:
    """Normalize cell values so JSON and SQLite receive predictable values."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, str):
        value = re.sub(r"\s+", " ", value).strip()
        return value or None
    return value


def normalize_coordinate(value: Any, field: str) -> Any:
    """Convert only latitude/longitude values entered in microdegrees."""
    limits = {"latitude": 90, "longitude": 180}
    limit = limits.get(field)
    if limit and isinstance(value, (int, float)) and limit < abs(value) < 100_000_000:
        return value / 1_000_000
    return value


def unique_headers(values: Iterable[Any]) -> list[str]:
    """Create non-empty, non-duplicated names for a worksheet header row."""
    used: Counter[str] = Counter()
    result = []
    for index, value in enumerate(values, 1):
        base = slug(str(value or "")) or f"column_{index}"
        used[base] += 1
        result.append(base if used[base] == 1 else f"{base}_{used[base]}")
    return result


def records_from_section(sheet: Any, header_row: int, end_row: int | None = None) -> list[dict[str, Any]]:
    """Transform one contiguous Excel table into records with source lineage."""
    # Normalize headers once, then apply them to every data row in the section.
    headers = unique_headers(cell.value for cell in sheet[header_row])
    output = []
    for excel_row, cells in enumerate(sheet.iter_rows(min_row=header_row + 1, max_row=end_row, values_only=True), header_row + 1):
        # Coordinate conversion is field-aware; ordinary prices must never be scaled.
        values = [normalize_coordinate(json_value(value), header) for header, value in zip(headers, cells)]
        # Do not create a record for visual spacer rows in the source workbook.
        if not any(value is not None for value in values):
            continue
        # Omit absent optional cells while retaining sheet and row traceability.
        record = {header: value for header, value in zip(headers, values) if value is not None}
        record["_source_sheet"] = sheet.title
        record["_source_row"] = excel_row
        output.append(record)
    return output


def extract_workbook(path: Path) -> dict[str, list[dict[str, Any]]]:
    """Extract every supported worksheet into named, normalized record sets."""
    # Read calculated values only and stream rows to keep memory use low.
    workbook = load_workbook(path, data_only=True, read_only=True)
    tables: dict[str, list[dict[str, Any]]] = {}
    for sheet in workbook.worksheets:
        if sheet.title == "REGRAS_CURADORIA":
            # This sheet contains two legitimate tables with different schemas.
            tables["curation_principles"] = records_from_section(sheet, 4, 10)
            tables["curation_rules"] = records_from_section(sheet, 14, 65)
        else:
            # Fail explicitly when a new sheet has no declared data-table header.
            header_row = HEADER_ROWS.get(sheet.title)
            if header_row is None:
                raise ValueError(f"No header mapping configured for worksheet {sheet.title!r}")
            tables[slug(sheet.title)] = records_from_section(sheet, header_row)
    # Business terminology overrides are versioned separately from the original workbook.
    overrides_path = Path(__file__).resolve().parents[1] / "config" / "curation_rule_overrides.json"
    overrides = json.loads(overrides_path.read_text(encoding="utf-8"))
    for record in tables.get("curation_rules", []):
        record.update(overrides.get(record.get("id"), {}))
    workbook.close()
    return tables


def document_hashes(path: Path) -> tuple[str, str]:
    """Return a file-instance ID and a separate hash for duplicate-content analysis."""
    content = path.read_bytes()
    content_hash = hashlib.sha256(content).hexdigest()
    # The filename distinguishes separately supplied copies with identical contents.
    instance_id = hashlib.sha256(path.as_posix().encode("utf-8") + b"\0" + content).hexdigest()
    return instance_id, content_hash


def extract_pdfs(directory: Path, paths: list[Path] | None = None) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    """Extract document-level metadata and page-level text from every PDF."""
    documents, pages, errors = [], [], []
    for path in sorted(paths if paths is not None else directory.rglob("*.pdf")):
        # Keep the file-instance ID as the shared primary key between document/pages.
        digest, content_hash = document_hashes(path)
        try:
            reader = PdfReader(str(path))
            page_texts = []
            for page_number, page in enumerate(reader.pages, 1):
                # Collapse layout whitespace while preserving all readable page text.
                text = re.sub(r"\s+", " ", page.extract_text() or "").strip()
                page_texts.append(text)
                pages.append({
                    "document_id": digest, "page_number": page_number, "text": text,
                    "source_file": path.as_posix(),
                })
            full_text = "\n".join(page_texts)
            metadata = reader.metadata or {}
            # Store source fields and derived search facets in one document record.
            documents.append({
                "document_id": digest,
                "content_sha256": content_hash,
                "file_name": path.name,
                "source_file": path.as_posix(),
                "file_size_bytes": path.stat().st_size,
                "page_count": len(reader.pages),
                "text_characters": len(full_text),
                "title": str(metadata.get("/Title") or "").strip() or None,
                "emails": sorted(set(EMAIL.findall(full_text))),
                "phones": sorted(set(re.sub(r"\s+", " ", item).strip() for item in PHONE.findall(full_text))),
                "urls": sorted(set(URL.findall(full_text))),
                "price_mentions": sorted(set(MONEY.findall(full_text))),
                "text": full_text,
            })
        except Exception as exc:  # one damaged PDF must not prevent the rest loading
            errors.append({"source_file": path.as_posix(), "error": f"{type(exc).__name__}: {exc}"})
    return documents, pages, errors


def write_jsonl(path: Path, records: list[dict[str, Any]]) -> None:
    """Write a record set as UTF-8 JSON Lines for streaming consumers."""
    with path.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False, default=str) + "\n")


def sqlite_type(value: Any) -> str:
    """Choose the narrowest SQLite affinity required by a sample value."""
    if isinstance(value, bool) or isinstance(value, int):
        return "INTEGER"
    if isinstance(value, float):
        return "REAL"
    return "TEXT"


def load_table(connection: sqlite3.Connection, name: str, records: list[dict[str, Any]]) -> None:
    """Create and populate one SQLite table from records with a flexible schema."""
    if not records:
        return
    # Preserve first-seen key order while allowing sparse records to add columns.
    columns = list(dict.fromkeys(key for record in records for key in record))
    types = {column: next((sqlite_type(record[column]) for record in records if record.get(column) is not None), "TEXT") for column in columns}
    # Quote generated identifiers before composing DDL from source-derived headers.
    quoted = lambda identifier: '"' + identifier.replace('"', '""') + '"'
    connection.execute(f"CREATE TABLE {quoted(name)} ({', '.join(f'{quoted(c)} {types[c]}' for c in columns)})")
    placeholders = ", ".join("?" for _ in columns)
    values = []
    for record in records:
        # Serialize nested PDF facets; SQLite scalar columns hold all other values.
        values.append(tuple(json.dumps(record.get(c), ensure_ascii=False) if isinstance(record.get(c), (list, dict)) else record.get(c) for c in columns))
    connection.executemany(f"INSERT INTO {quoted(name)} VALUES ({placeholders})", values)


def _build(data_dir: Path, output_dir: Path) -> dict[str, Any]:
    """Build a delivery in an isolated staging directory."""
    workbook, supplement, pdf_paths, source_manifest = LocalSource(data_dir).discover()
    # Extract both source types, then add their records to one logical data catalog.
    tables = extract_workbook(workbook)
    # Preserve other workbook schemas as source evidence, without guessing mappings.
    reference_rows = []
    for item in source_manifest["objects"]:
        if item["role"] != "reference_workbook":
            continue
        reference_path = data_dir / item["key"]
        reference = load_workbook(reference_path, read_only=True, data_only=True)
        try:
            for sheet in reference:
                for number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                    if any(value is not None for value in cells):
                        reference_rows.append({"_source_file": reference_path.as_posix(), "_source_sheet": sheet.title, "_source_row": number, "cells": [value.isoformat() if isinstance(value, datetime) else value for value in cells]})
        finally:
            reference.close()
    tables["source_workbook_rows"] = reference_rows
    structured = None
    if supplement is not None:
        from etl.structured import integrate_structured
        structured = integrate_structured(supplement, tables)
    documents, pages, pdf_errors = extract_pdfs(data_dir, pdf_paths)
    tables["proposal_documents"] = documents
    tables["proposal_pages"] = pages

    from etl.portal_catalog import integrate_portal
    source_manifest["objects"].append(integrate_portal(data_dir.resolve(), tables))

    # Validate transformed data before publishing any metrics or analytical models.
    validation_report, validation_issues = validate(tables)
    # Only validated records feed the descriptive KPIs and transparent score models.
    kpis, models, model_outputs = build_analytics(tables, validation_report)
    if validation_report["status"] != "passed":
        raise ValueError(f"Warehouse validation failed: {validation_report['issue_count']} issues")
    output_dir.mkdir(parents=True, exist_ok=True)
    for table, records in tables.items():
        # JSONL is the portable, table-per-file delivery format.
        write_jsonl(output_dir / f"{table}.jsonl", records)
    write_jsonl(output_dir / "data_quality_issues.jsonl", validation_issues)
    for model_name, records in model_outputs.items():
        write_jsonl(output_dir / f"{model_name}.jsonl", records)
    (output_dir / "validation_report.json").write_text(json.dumps(validation_report, ensure_ascii=False, indent=2), encoding="utf-8")
    (output_dir / "kpis.json").write_text(json.dumps(kpis, ensure_ascii=False, indent=2), encoding="utf-8")
    (output_dir / "analytical_models.json").write_text(json.dumps(models, ensure_ascii=False, indent=2), encoding="utf-8")

    # SQLite is the relational delivery format for SQL queries and joins.
    database = output_dir / "blu_etl.sqlite"
    with closing(sqlite3.connect(database)) as connection:
        for table, records in tables.items():
            load_table(connection, table, records)
        load_table(connection, "data_quality_issues", validation_issues)
        for model_name, records in model_outputs.items():
            load_table(connection, model_name, records)
        if pages:
            connection.execute("CREATE INDEX idx_proposal_pages_document ON proposal_pages(document_id)")
        connection.commit()

    # Publish run metadata so consumers can verify data coverage and quality.
    manifest = {
        "pipeline": "blu-local-etl", "ran_at_utc": datetime.now(UTC).isoformat(),
        "input_directory": data_dir.as_posix(), "workbook": workbook.as_posix(), "source": source_manifest,
        "tables": {table: len(records) for table, records in tables.items()},
        "pdfs_discovered": len(pdf_paths),
        "pdf_errors": pdf_errors,
        "structured_dataset": structured,
        "validation": {"status": validation_report["status"], "issue_count": validation_report["issue_count"]},
        "analytics": {"kpis": "kpis.json", "models": "analytical_models.json"},
        "outputs": {"database": database.name, "format": "JSON Lines + SQLite"},
    }
    (output_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest



def run(data_dir: Path, output_dir: Path) -> dict[str, Any]:
    """Read local raw data, validate a complete delivery, then publish with rollback."""
    source, target = data_dir.resolve(), output_dir.resolve()
    if source == target or source in target.parents or target in source.parents:
        raise ValueError("Source and warehouse directories must be separate and not nested")
    if output_dir.is_symlink():
        raise ValueError("Warehouse must not be a symbolic link")
    target.parent.mkdir(parents=True, exist_ok=True)
    stage = target.with_name(f".{target.name}-staging-{uuid.uuid4().hex}")
    stage.mkdir()  # Inherit warehouse-parent permissions; tempfile uses restrictive ACLs on Windows.
    backup = target.with_name(f".{target.name}-backup-{uuid.uuid4().hex}")
    # Every directory removed or moved below is a checked sibling of the explicit output.
    assert stage.parent == target.parent and backup.parent == target.parent
    try:
        manifest = _build(data_dir, stage)
        if target.exists():
            target.rename(backup)
        try:
            stage.rename(target)
        except Exception:
            if backup.exists():
                backup.rename(target)
            raise
        if backup.exists():
            shutil.rmtree(backup, ignore_errors=True)
        return manifest
    finally:
        if stage.exists():
            shutil.rmtree(stage)


def main() -> None:
    """Expose the ETL as ``python -m etl.pipeline``."""
    parser = argparse.ArgumentParser(description="Run the BLU local data ETL")
    parser.add_argument("--data-dir", type=Path, default=Path("data"))
    parser.add_argument("--output-dir", type=Path, default=Path("warehouse"))
    args = parser.parse_args()
    manifest = run(args.data_dir, args.output_dir)
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    # Run the command-line entry point only when this module is executed directly.
    main()
