"""Cria fontes sintéticas mínimas para validar o ETL em CI sem dados privados."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook


def create(destination: Path) -> Path:
    """Escreve um catálogo pequeno, válido e determinístico no destino indicado."""
    destination.mkdir(parents=True, exist_ok=True)
    target = destination / "ci-catalog.xlsx"
    workbook = Workbook()

    attractions = workbook.active
    attractions.title = "ATRACOES"
    attractions.append(["Fixture sintética de CI"])
    for _ in range(4):
        attractions.append([])
    attractions.append([
        "ID",
        "Cidade",
        "Nome da atração",
        "Categoria",
        "Descrição curada",
        "Latitude",
        "Longitude",
    ])
    attractions.append([
        "CI-ATR-001",
        "Lisboa",
        "Atração sintética",
        "Cultura",
        "Registo fictício usado apenas para validar o pipeline.",
        38.7223,
        -9.1393,
    ])

    restaurants = workbook.create_sheet("RESTAURANTES")
    restaurants.append(["ID BLU", "Estabelecimento", "Cidade"])
    restaurants.append(["CI-RES-001", "Restaurante sintético", "Lisboa"])

    experiences = workbook.create_sheet("EXPERIENCIAS")
    experiences.append(["Nome da experiência", "Categoria", "Localização"])
    experiences.append(["Experiência sintética", "Cultura", "Lisboa"])

    prices = workbook.create_sheet("PREÇOS")
    prices.append(["ID", "Categoria", "Preço referência EUR"])
    prices.append(["CI-PRE-001", "Cultura", 10])

    requests = workbook.create_sheet("DATA_REQUEST")
    for _ in range(5):
        requests.append([])
    requests.append(["ID", "Estado"])
    requests.append(["CI-REQ-001", "Sintético"])

    rules = workbook.create_sheet("REGRAS_CURADORIA")
    for _ in range(3):
        rules.append([])
    rules.append(["ID", "Princípio"])
    rules.append(["CI-PRI-001", "Revisão humana"])
    for _ in range(8):
        rules.append([])
    rules.append(["ID", "Regra"])
    rules.append(["CI-REG-001", "Não inferir dados desconhecidos"])

    workbook.save(target)
    workbook.close()
    return target


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    print(create(args.output))


if __name__ == "__main__":
    main()
